"""
Excel File Upload and Data Filler Script
Reads an Excel file, finds empty cells, and fills them with data from a source
"""

import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd
from typing import Dict, List, Tuple
import sys

class ExcelDataFiller:
    def __init__(self, file_path: str):
        """Initialize with Excel file path"""
        self.file_path = file_path
        self.workbook = None
        self.worksheet = None
        
    def load_excel(self) -> bool:
        """Load Excel file"""
        try:
            self.workbook = openpyxl.load_workbook(self.file_path)
            self.worksheet = self.workbook.active
            print(f"✓ Successfully loaded Excel file: {self.file_path}")
            return True
        except Exception as e:
            print(f"✗ Error loading Excel file: {e}")
            return False
    
    def find_empty_cells(self) -> List[Tuple[int, int]]:
        """Find all empty cells in the worksheet"""
        empty_cells = []
        for row in self.worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    empty_cells.append((cell.row, cell.column))
        
        print(f"✓ Found {len(empty_cells)} empty cells")
        return empty_cells
    
    def find_empty_cells_by_column(self, column: str) -> List[Tuple[int, str]]:
        """Find empty cells in a specific column"""
        empty_cells = []
        col_index = openpyxl.utils.column_index_from_string(column)
        
        for row in range(2, self.worksheet.max_row + 1):  # Start from row 2 (skip header)
            cell = self.worksheet.cell(row=row, column=col_index)
            if cell.value is None:
                empty_cells.append((row, column))
        
        print(f"✓ Found {len(empty_cells)} empty cells in column {column}")
        return empty_cells
    
    def fill_empty_cells(self, data_source: Dict[Tuple[int, int], str]) -> bool:
        """
        Fill empty cells with data from source
        data_source: Dictionary with (row, col) as key and value to fill
        """
        try:
            for (row, col), value in data_source.items():
                cell = self.worksheet.cell(row=row, column=col)
                cell.value = value
                print(f"  Filled cell ({row}, {col}) with: {value}")
            
            print(f"✓ Successfully filled {len(data_source)} cells")
            return True
        except Exception as e:
            print(f"✗ Error filling cells: {e}")
            return False
    
    def save_excel(self, output_path: str = None) -> bool:
        """Save the modified Excel file"""
        try:
            save_path = output_path or self.file_path
            self.workbook.save(save_path)
            print(f"✓ Successfully saved Excel file: {save_path}")
            return True
        except Exception as e:
            print(f"✗ Error saving Excel file: {e}")
            return False
    
    def get_empty_data_map(self, column: str, lookup_func) -> Dict[Tuple[int, str], str]:
        """
        Get a map of empty cells and their looked-up values
        lookup_func: Function that takes (row_number, row_data) and returns filled value
        """
        col_index = openpyxl.utils.column_index_from_string(column)
        data_map = {}
        
        for row in range(2, self.worksheet.max_row + 1):
            cell = self.worksheet.cell(row=row, column=col_index)
            if cell.value is None:
                # Get row data for context
                row_data = {}
                for col_num in range(1, self.worksheet.max_column + 1):
                    col_letter = get_column_letter(col_num)
                    row_data[col_letter] = self.worksheet.cell(row=row, column=col_num).value
                
                # Call lookup function
                filled_value = lookup_func(row, row_data)
                if filled_value:
                    data_map[(row, col_index)] = filled_value
        
        return data_map


def lookup_from_google_sheets(sheet_id: str, range_name: str):
    """
    Lookup data from Google Sheets
    Requires: pip install google-auth-oauthlib google-auth-httplib2 google-api-python-client
    """
    try:
        from google.auth.transport.requests import Request
        from google.oauth2.service_account import Credentials
        from google_auth_oauthlib.flow import InstalledAppFlow
        from googleapiclient.discovery import build
        
        SCOPES = ['https://www.googleapis.com/auth/spreadsheets.readonly']
        
        # This requires a credentials.json file
        service = build('sheets', 'v4', credentials=get_google_credentials(SCOPES))
        sheet = service.spreadsheets().values().get(
            spreadsheetId=sheet_id,
            range=range_name
        ).execute()
        
        return sheet.get('values', [])
    except ImportError:
        print("✗ Google API libraries not installed. Install with:")
        print("  pip install google-auth-oauthlib google-auth-httplib2 google-api-python-client")
        return None


def get_google_credentials(scopes):
    """Get Google API credentials"""
    from google_auth_oauthlib.flow import InstalledAppFlow
    
    flow = InstalledAppFlow.from_client_secrets_file(
        'credentials.json', scopes)
    creds = flow.run_local_server(port=0)
    return creds


# Example usage function
def example_lookup_function(row_num: int, row_data: dict) -> str:
    """
    Example lookup function - customize this based on your data source
    Returns the value to fill in the empty cell
    """
    # Example: Look up based on another column
    # e.g., if 'Name' exists, look it up from Google
    
    name = row_data.get('A')  # Assuming name is in column A
    if name:
        # Here you would call your actual data source
        # For demo: return a dummy value
        return f"Data for {name}"
    return None


if __name__ == "__main__":
    # Configuration
    EXCEL_FILE = "data.xlsx"  # Change to your Excel file path
    OUTPUT_FILE = "data_filled.xlsx"
    TARGET_COLUMN = "B"  # Change to the column with empty cells
    
    # Initialize filler
    filler = ExcelDataFiller(EXCEL_FILE)
    
    # Load Excel file
    if not filler.load_excel():
        sys.exit(1)
    
    # Find empty cells in target column
    empty_cells = filler.find_empty_cells_by_column(TARGET_COLUMN)
    
    if not empty_cells:
        print("✓ No empty cells found!")
        sys.exit(0)
    
    # Get data to fill (using example function)
    print("\n📊 Looking up data for empty cells...")
    data_to_fill = filler.get_empty_data_map(TARGET_COLUMN, example_lookup_function)
    
    if not data_to_fill:
        print("✗ No data found to fill")
        sys.exit(1)
    
    # Fill the cells
    print("\n✏️  Filling empty cells...")
    col_index = openpyxl.utils.column_index_from_string(TARGET_COLUMN)
    fill_map = {(row, col_index): val for (row, col_index), val in data_to_fill.items()}
    
    if filler.fill_empty_cells(fill_map):
        # Save the updated file
        print(f"\n💾 Saving updated Excel file...")
        filler.save_excel(OUTPUT_FILE)
        print(f"\n✅ Complete! Saved to: {OUTPUT_FILE}")
    else:
        print("✗ Error filling cells")
        sys.exit(1)
